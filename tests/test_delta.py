import json
from pathlib import Path

import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from sanctions_parser.config import SourceConfig
from sanctions_parser.delta import (
    build_snapshot,
    compare_snapshots,
    export_delta_report,
    load_baseline_ref,
    load_snapshot,
    save_baseline,
)
from sanctions_parser.downloader import DownloadResult, sha256_file
from sanctions_parser.models import Address, Alias, Entity, NormalizedData
from sanctions_parser.pipeline import (
    PARSER_SCHEMA_VERSION,
    find_existing_raw_baseline,
    run_delta_source,
)


def snapshot(
    data: NormalizedData,
    checksum: str,
    created_at: str,
    raw_path: Path | None = None,
):
    return build_snapshot(
        data,
        source="ofac",
        checksum=checksum,
        parser_schema_version=PARSER_SCHEMA_VERSION,
        raw_path=raw_path or Path(f"{checksum}.xml"),
        created_at=created_at,
    )


def test_first_comparison_creates_baseline_without_marking_records_new() -> None:
    current = snapshot(
        NormalizedData(entities=[Entity("1", "ofac", "Entity", "Example")]),
        "current",
        "2026-08-01T00:00:00+00:00",
    )

    result = compare_snapshots(None, current)

    assert result.summary.result == "baseline_created"
    assert result.summary.new == 0
    assert result.summary.unchanged == 1
    assert result.changes == ()


def test_existing_raw_baseline_chooses_latest_archived_xml(tmp_path: Path) -> None:
    old_raw = tmp_path / "raw" / "ofac" / "2026-07-20" / "sdn.xml"
    old_raw.parent.mkdir(parents=True)
    old_raw.write_text("<old/>", encoding="utf-8")
    latest_raw = tmp_path / "raw" / "ofac" / "2026-07-26" / "sdn.XML"
    latest_raw.parent.mkdir(parents=True)
    latest_raw.write_text("<latest/>", encoding="utf-8")
    (latest_raw.parent / "notes.txt").write_text("ignore", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")

    candidate = find_existing_raw_baseline(source, tmp_path)

    assert candidate is not None
    assert candidate.path == latest_raw
    assert candidate.archive_date == "2026-07-26"


def test_existing_raw_is_imported_before_latest_download_and_compared(
    monkeypatch, tmp_path: Path
) -> None:
    old_raw = tmp_path / "raw" / "ofac" / "2026-07-26" / "sdn.xml"
    old_raw.parent.mkdir(parents=True)
    old_raw.write_text("<old/>", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    candidate = find_existing_raw_baseline(source, tmp_path)
    assert candidate is not None
    current_raw = tmp_path / "raw" / "ofac" / "2026-08-07" / "sdn.xml"
    current_raw.parent.mkdir(parents=True)
    current_raw.write_text("<current/>", encoding="utf-8")
    events: list[str] = []

    def fake_parse(_name, path):
        events.append(f"parse:{path.parent.name}")
        primary_name = "Old Name" if path == old_raw else "Current Name"
        return NormalizedData(
            entities=[Entity("1", "ofac", "Entity", primary_name)]
        )

    def fake_download(*args, **kwargs):
        events.append("download")
        return DownloadResult("ofac", current_raw, "current-checksum", True, 10)

    monkeypatch.setattr("sanctions_parser.pipeline.parse", fake_parse)
    monkeypatch.setattr("sanctions_parser.pipeline.download_source", fake_download)

    result = run_delta_source(
        source,
        tmp_path,
        show_download_progress=False,
        initial_baseline=candidate,
    )
    baseline = load_baseline_ref(tmp_path, "ofac")

    assert events == ["parse:2026-07-26", "download", "parse:2026-08-07"]
    assert result.status == "changes"
    assert result.summary is not None
    assert result.summary.updated == 1
    assert result.summary.previous_checksum == sha256_file(old_raw)
    assert result.summary.current_checksum == "current-checksum"
    assert baseline is not None and baseline.checksum == "current-checksum"


def test_failed_existing_raw_import_does_not_download_or_create_baseline(
    monkeypatch, tmp_path: Path
) -> None:
    old_raw = tmp_path / "raw" / "ofac" / "2026-07-26" / "sdn.xml"
    old_raw.parent.mkdir(parents=True)
    old_raw.write_text("<broken/>", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    candidate = find_existing_raw_baseline(source, tmp_path)
    assert candidate is not None
    download_called = False

    monkeypatch.setattr(
        "sanctions_parser.pipeline.parse",
        lambda *args, **kwargs: NormalizedData(),
    )

    def fake_download(*args, **kwargs):
        nonlocal download_called
        download_called = True
        raise AssertionError("download must not run after invalid baseline")

    monkeypatch.setattr("sanctions_parser.pipeline.download_source", fake_download)

    with pytest.raises(ValueError, match="produced zero entities"):
        run_delta_source(
            source,
            tmp_path,
            show_download_progress=False,
            initial_baseline=candidate,
        )

    assert download_called is False
    assert load_baseline_ref(tmp_path, "ofac") is None


def test_entity_and_name_alias_changes_are_compared_by_stable_id() -> None:
    previous = snapshot(
        NormalizedData(
            entities=[
                Entity("1", "ofac", "Entity", "Old Company"),
                Entity("2", "ofac", "Individual", "Removed Person"),
            ],
            aliases=[
                Alias("1", "Shared Alias", "Low"),
                Alias("1", "Removed Alias", "Good"),
                Alias("2", "Old Nickname", "Low"),
            ],
        ),
        "old",
        "2026-08-01T00:00:00+00:00",
    )
    current = snapshot(
        NormalizedData(
            entities=[
                Entity("1", "ofac", "Entity", "New Company"),
                Entity("3", "ofac", "Individual", "New Person"),
            ],
            aliases=[
                Alias("1", "Shared Alias", "Good"),
                Alias("1", "Added Alias", "f.k.a."),
                Alias("3", "New Nickname", "Good"),
            ],
        ),
        "new",
        "2026-08-08T00:00:00+00:00",
    )

    result = compare_snapshots(previous, current)

    assert result.summary.new == 1
    assert result.summary.updated == 1
    assert result.summary.removed == 1
    assert result.summary.unchanged == 0
    rows = {row["entity_id"]: row for row in result.changes}
    assert rows["1"]["status"] == "UPDATED"
    assert "Primary name changed" in rows["1"]["what_changed"]
    assert "alias added" in rows["1"]["what_changed"]
    assert "alias removed" in rows["1"]["what_changed"]
    assert "alias strength change" in rows["1"]["what_changed"]
    assert rows["1"]["previous_primary_name"] == "Old Company"
    assert rows["2"]["status"] == "REMOVED"
    assert rows["2"]["primary_name"] == "Removed Person"
    assert rows["3"]["status"] == "NEW"


def test_alias_order_duplicates_and_spacing_do_not_create_false_changes() -> None:
    previous = snapshot(
        NormalizedData(
            entities=[Entity("1", "ofac", "Entity", "Example Company")],
            aliases=[
                Alias("1", "Alpha Alias", "Good"),
                Alias("1", "Beta Alias", "Low"),
            ],
        ),
        "old",
        "2026-08-01T00:00:00+00:00",
    )
    current = snapshot(
        NormalizedData(
            entities=[Entity("1", "ofac", "Entity", " Example   Company ")],
            aliases=[
                Alias("1", "Beta Alias", "weak"),
                Alias("1", "Alpha   Alias", "strong"),
                Alias("1", "Alpha Alias", "Good"),
            ],
        ),
        "new",
        "2026-08-08T00:00:00+00:00",
    )

    result = compare_snapshots(previous, current)

    assert result.summary.result == "no_changes"
    assert result.summary.unchanged == 1
    assert result.changes == ()


def test_non_name_relation_change_is_reported_with_name_and_aliases() -> None:
    previous = snapshot(
        NormalizedData(
            entities=[Entity("1", "ofac", "Entity", "Example")],
            aliases=[Alias("1", "Example Alias", "Good")],
            addresses=[Address("1", "Old Street")],
        ),
        "old",
        "2026-08-01T00:00:00+00:00",
    )
    current = snapshot(
        NormalizedData(
            entities=[Entity("1", "ofac", "Entity", "Example")],
            aliases=[Alias("1", "Example Alias", "Good")],
            addresses=[Address("1", "New Street")],
        ),
        "new",
        "2026-08-08T00:00:00+00:00",
    )

    result = compare_snapshots(previous, current)

    assert result.summary.updated == 1
    assert result.changes[0]["primary_name"] == "Example"
    assert result.changes[0]["aliases"] == "Example Alias [strong]"
    assert result.changes[0]["what_changed"] == "Addresses changed"


def test_non_latin_alias_change_is_preserved() -> None:
    previous = snapshot(
        NormalizedData(
            entities=[Entity("1", "ofac", "Individual", "Example Person")],
            aliases=[Alias("1", "محمد", "Good", "Arabic")],
        ),
        "old",
        "2026-08-01T00:00:00+00:00",
    )
    current = snapshot(
        NormalizedData(
            entities=[Entity("1", "ofac", "Individual", "Example Person")],
            aliases=[Alias("1", "محمد علي", "Good", "Arabic")],
        ),
        "new",
        "2026-08-08T00:00:00+00:00",
    )

    result = compare_snapshots(previous, current)

    assert result.summary.updated == 1
    assert "محمد علي [strong]" in result.changes[0]["aliases_added"]
    assert "محمد [strong]" in result.changes[0]["aliases_removed"]


def test_delta_report_exports_csv_excel_and_parquet(tmp_path: Path) -> None:
    previous = snapshot(
        NormalizedData(entities=[Entity("1", "ofac", "Entity", "Old")]),
        "old",
        "2026-08-01T00:00:00+00:00",
    )
    current = snapshot(
        NormalizedData(entities=[Entity("1", "ofac", "Entity", "New")]),
        "new",
        "2026-08-08T00:00:00+00:00",
    )
    comparison = compare_snapshots(previous, current)

    export_delta_report(comparison, tmp_path)

    assert (tmp_path / "csv" / "change_summary.csv").is_file()
    assert (tmp_path / "csv" / "name_alias_changes.csv").is_file()
    assert (
        pq.ParquetFile(
            tmp_path / "parquet" / "name_alias_changes.parquet"
        ).metadata.num_rows
        == 1
    )
    workbook = load_workbook(
        tmp_path / "excel" / "delta-report.xlsx",
        read_only=False,
        data_only=True,
    )
    try:
        assert workbook.sheetnames[:2] == ["change summary", "name+alias changes"]
        sheet = workbook["name+alias changes"]
        assert sheet["A2"].value == "UPDATED"
        assert sheet["A2"].fill.fgColor.rgb == "00FCE4D6"
        assert sheet.freeze_panes == "A2"
    finally:
        workbook.close()


def test_successful_delta_advances_baseline_only_after_report(
    monkeypatch, tmp_path: Path
) -> None:
    first_raw = tmp_path / "raw" / "ofac" / "first.xml"
    first_raw.parent.mkdir(parents=True)
    first_raw.write_text("<sdnList/>", encoding="utf-8")
    second_raw = tmp_path / "raw" / "ofac" / "second.xml"
    second_raw.write_text("<sdnList/>", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    downloads = iter(
        [
            DownloadResult("ofac", first_raw, "first", True, 10),
            DownloadResult("ofac", second_raw, "second", True, 10),
        ]
    )
    parsed = iter(
        [
            NormalizedData(entities=[Entity("1", "ofac", "Entity", "Old Name")]),
            NormalizedData(entities=[Entity("1", "ofac", "Entity", "New Name")]),
        ]
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.download_source",
        lambda *args, **kwargs: next(downloads),
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.parse",
        lambda *args, **kwargs: next(parsed),
    )

    first = run_delta_source(source, tmp_path, show_download_progress=False)
    first_baseline = load_baseline_ref(tmp_path, "ofac")
    second = run_delta_source(source, tmp_path, show_download_progress=False)
    second_baseline = load_baseline_ref(tmp_path, "ofac")

    assert first.status == "baseline_created"
    assert first_baseline is not None and first_baseline.checksum == "first"
    assert second.status == "changes"
    assert second.summary is not None and second.summary.updated == 1
    assert second_baseline is not None and second_baseline.checksum == "second"
    assert load_snapshot(second_baseline.snapshot_path).checksum == "second"


def test_failed_delta_export_preserves_previous_baseline(
    monkeypatch, tmp_path: Path
) -> None:
    raw = tmp_path / "raw" / "ofac" / "current.xml"
    raw.parent.mkdir(parents=True)
    raw.write_text("<sdnList/>", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    downloads = iter(
        [
            DownloadResult("ofac", raw, "first", True, 10),
            DownloadResult("ofac", raw, "second", True, 10),
        ]
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.download_source",
        lambda *args, **kwargs: next(downloads),
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.parse",
        lambda *args, **kwargs: NormalizedData(
            entities=[Entity("1", "ofac", "Entity", "Example")]
        ),
    )

    run_delta_source(source, tmp_path, show_download_progress=False)
    baseline_before = json.loads(
        (tmp_path / ".state" / "delta" / "ofac.json").read_text(encoding="utf-8")
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.export_delta_report",
        lambda *args, **kwargs: (_ for _ in ()).throw(OSError("disk full")),
    )

    with pytest.raises(OSError, match="disk full"):
        run_delta_source(source, tmp_path, show_download_progress=False)

    baseline_after = json.loads(
        (tmp_path / ".state" / "delta" / "ofac.json").read_text(encoding="utf-8")
    )
    assert baseline_after == baseline_before


def test_delta_after_ordinary_download_still_uses_older_delta_checkpoint(
    monkeypatch, tmp_path: Path
) -> None:
    old_raw = tmp_path / "raw" / "ofac" / "old.xml"
    old_raw.parent.mkdir(parents=True)
    old_raw.write_text("<sdnList/>", encoding="utf-8")
    current_raw = tmp_path / "raw" / "ofac" / "current.xml"
    current_raw.write_text("<sdnList/>", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    old_snapshot = snapshot(
        NormalizedData(entities=[Entity("1", "ofac", "Entity", "Old")]),
        "old-checksum",
        "2026-08-01T00:00:00+00:00",
        old_raw,
    )
    save_baseline(
        tmp_path,
        old_snapshot,
        last_checked_at="2026-08-01T00:00:00+00:00",
    )
    state_root = tmp_path / ".state"
    state_root.mkdir(exist_ok=True)
    (state_root / "ofac.json").write_text(
        json.dumps(
            {
                "checksum": "current-checksum",
                "path": str(current_raw),
                "url": source.url,
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.download_source",
        lambda *args, **kwargs: DownloadResult(
            "ofac", None, "current-checksum", False, 10
        ),
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.parse",
        lambda *args, **kwargs: NormalizedData(
            entities=[Entity("1", "ofac", "Entity", "Current")]
        ),
    )

    result = run_delta_source(source, tmp_path, show_download_progress=False)

    assert result.summary is not None
    assert result.summary.updated == 1
    assert result.summary.previous_checksum == "old-checksum"
    assert result.summary.current_checksum == "current-checksum"


def test_parser_version_change_rebuilds_old_checkpoint_with_current_rules(
    monkeypatch, tmp_path: Path
) -> None:
    raw = tmp_path / "raw" / "ofac" / "list.xml"
    raw.parent.mkdir(parents=True)
    raw.write_text("<sdnList/>", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    old_snapshot = build_snapshot(
        NormalizedData(entities=[Entity("1", "ofac", "Entity", "Old parser name")]),
        source="ofac",
        checksum="same-checksum",
        parser_schema_version=PARSER_SCHEMA_VERSION - 1,
        raw_path=raw,
        created_at="2026-08-01T00:00:00+00:00",
    )
    save_baseline(
        tmp_path,
        old_snapshot,
        last_checked_at="2026-08-01T00:00:00+00:00",
    )
    monkeypatch.setattr(
        "sanctions_parser.pipeline.download_source",
        lambda *args, **kwargs: DownloadResult("ofac", raw, "same-checksum", True, 10),
    )
    parse_calls: list[Path] = []

    def current_parser(_name, path):
        parse_calls.append(path)
        return NormalizedData(
            entities=[Entity("1", "ofac", "Entity", "Current parser name")]
        )

    monkeypatch.setattr("sanctions_parser.pipeline.parse", current_parser)

    result = run_delta_source(source, tmp_path, show_download_progress=False)
    baseline = load_baseline_ref(tmp_path, "ofac")

    assert result.status == "no_changes"
    assert len(parse_calls) == 1
    assert baseline is not None
    assert baseline.parser_schema_version == PARSER_SCHEMA_VERSION


def test_corrupt_delta_state_is_not_silently_replaced(
    monkeypatch, tmp_path: Path
) -> None:
    raw = tmp_path / "raw" / "ofac" / "list.xml"
    raw.parent.mkdir(parents=True)
    raw.write_text("<sdnList/>", encoding="utf-8")
    state_path = tmp_path / ".state" / "delta" / "ofac.json"
    state_path.parent.mkdir(parents=True)
    state_path.write_text("{broken", encoding="utf-8")
    source = SourceConfig("ofac", "https://example.test/ofac.xml", "ofac")
    monkeypatch.setattr(
        "sanctions_parser.pipeline.download_source",
        lambda *args, **kwargs: DownloadResult("ofac", raw, "new", True, 10),
    )

    with pytest.raises(ValueError, match="Invalid OFAC delta baseline state"):
        run_delta_source(source, tmp_path, show_download_progress=False)

    assert state_path.read_text(encoding="utf-8") == "{broken"
