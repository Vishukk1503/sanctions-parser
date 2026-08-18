from pathlib import Path

import pyarrow.parquet as pq
from openpyxl import load_workbook

from sanctions_parser.exporter import NAME_ALIAS_COLUMNS, export_data
from sanctions_parser.models import Alias, Entity, NormalizedData


def test_empty_relations_are_exported_with_stable_schemas(tmp_path: Path) -> None:
    data = NormalizedData(entities=[Entity("1", "ofac", "Entity", "Test entity")])

    export_data(data, tmp_path)

    assert (tmp_path / "csv" / "alias.csv").read_text(encoding="utf-8").strip() == (
        "entity_id,alias,quality,language,first_name,middle_name,last_name"
    )
    assert pq.ParquetFile(tmp_path / "parquet" / "alias.parquet").metadata.num_rows == 0
    workbook = load_workbook(
        tmp_path / "excel" / "sanctions.xlsx", read_only=True, data_only=True
    )
    try:
        assert workbook["alias"].max_row == 1
        assert workbook["entity"].max_row == 2
        sheet = workbook["name+alias"]
        values = list(sheet.values)
        assert sheet.max_row == 2
        assert list(values[0]) == NAME_ALIAS_COLUMNS
        row = dict(zip(NAME_ALIAS_COLUMNS, values[1], strict=True))
        assert row["entity_id"] == "1"
        assert row["primary_name"] == "Test entity"
        assert row["PrimaryFullName"] == "Test entity"
        assert row["aliases"] is None
        assert row["alias_count"] == 0
    finally:
        workbook.close()


def test_excel_combines_multiple_aliases_by_entity_id(tmp_path: Path) -> None:
    data = NormalizedData(
        entities=[
            Entity(
                "1",
                "ofac",
                "Individual",
                "Primary Middle Person",
                primary_first_name="Primary",
                primary_middle_name="Middle",
                primary_last_name="Person",
            ),
            Entity("2", "ofac", "Entity", "Primary Company"),
        ],
        aliases=[
            Alias("1", "Weak Alias", "Low"),
            Alias("1", "Old Name", "f.k.a."),
            Alias(
                "1",
                "Strong Middle Alias",
                "Good",
                first_name="Strong",
                middle_name="Middle",
                last_name="Alias",
            ),
            Alias("1", "Regular Alias", "a.k.a."),
            Alias("1", "Strong Middle Alias", "weak"),
            Alias("1", "PRIMARY MIDDLE PERSON", "strong"),
            Alias("1", "اسم", "Good", "original script"),
            Alias("2", "Company Alias", "strong"),
        ],
    )

    export_data(data, tmp_path, export_csv=False, parquet=False)

    workbook = load_workbook(
        tmp_path / "excel" / "sanctions.xlsx", read_only=False, data_only=True
    )
    try:
        sheet = workbook["name+alias"]
        values = list(sheet.values)
        assert list(values[0]) == NAME_ALIAS_COLUMNS
        individual = dict(zip(NAME_ALIAS_COLUMNS, values[1], strict=True))
        organization = dict(zip(NAME_ALIAS_COLUMNS, values[2], strict=True))
        assert individual["primary_name"] == "Primary Middle Person"
        assert individual["PrimaryFirstName"] == "Primary"
        assert individual["PrimaryMiddleName"] == "Middle"
        assert individual["PrimaryLastName"] == "Person"
        assert individual["PrimaryIsBrokenName"] == "true"
        assert individual["Alias1FirstName"] == "Strong"
        assert individual["Alias1MiddleName"] == "Middle"
        assert individual["Alias1LastName"] == "Alias"
        assert individual["Alias1IsBrokenName"] == "true"
        assert individual["Alias2LastName"] == "Regular Alias"
        assert individual["Alias3LastName"] == "اسم"
        assert individual["aliases"] == (
            "Strong Middle Alias [strong] | Regular Alias [strong] | اسم [strong] | "
            "Weak Alias [weak] | Old Name [former]"
        )
        assert individual["alias_count"] == 5
        assert organization["PrimaryFullName"] == "Primary Company"
        assert organization["Alias1FullName"] == "Company Alias"
        assert organization["alias_count"] == 1
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:Y3"
        assert sheet["A1"].fill.fgColor.rgb == "000F4C81"
        assert sheet["A1"].font.bold is True
        assert sheet["A1"].font.color.rgb == "00FFFFFF"
        assert sheet["A2"].alignment.horizontal == "left"
        assert sheet["Y2"].number_format == "#,##0"
        assert sheet.tables["NameAliasTable"].ref == "A1:Y3"
    finally:
        workbook.close()

    assert not (tmp_path / "csv").exists()
    assert not (tmp_path / "parquet").exists()
    assert not (tmp_path / "excel" / "name+alias.csv").exists()
    assert not (tmp_path / "excel" / "name+alias.parquet").exists()
