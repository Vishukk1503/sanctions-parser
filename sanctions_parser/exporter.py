from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .models import NormalizedData
from .reporting import alias_tag
from .ssb import SSB_NAME_COLUMNS, export_ssb, ordered_aliases, project_name_fields

NAME_ALIAS_SHEET = "name+alias"
NAME_ALIAS_COLUMNS = [
    "entity_id",
    "record_type",
    "primary_name",
    *SSB_NAME_COLUMNS,
    "aliases",
    "alias_count",
]

def _name_alias_frame(data: NormalizedData) -> pd.DataFrame:
    """Return one row per entity with split names and every tagged alias."""
    aliases_by_entity = ordered_aliases(data)
    rows: list[dict[str, object]] = []
    for entity in data.entities:
        aliases = aliases_by_entity.get(entity.entity_id, [])
        rows.append(
            {
                "entity_id": entity.entity_id,
                "record_type": entity.record_type,
                "primary_name": entity.primary_name,
                **project_name_fields(entity, aliases),
                "aliases": " | ".join(
                    f"{item.alias} [{alias_tag(item.quality)}]" for item in aliases
                ),
                "alias_count": len(aliases),
            }
        )
    return pd.DataFrame(rows, columns=NAME_ALIAS_COLUMNS)


def _format_name_alias_sheet(sheet: Worksheet) -> None:
    """Apply compact, readable formatting to the convenience sheet."""
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(NAME_ALIAS_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="0F4C81")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for cell in sheet["A"][1:]:
        cell.alignment = Alignment(horizontal="left")

    for index, name in enumerate(NAME_ALIAS_COLUMNS, start=1):
        column = get_column_letter(index)
        if name == "aliases":
            width = 72
        elif name == "alias_count":
            width = 13
        elif name.endswith("FullName") or name == "primary_name":
            width = 38
        elif name.endswith("IsBrokenName"):
            width = 18
        else:
            width = 20
        sheet.column_dimensions[column].width = width

    for cell in sheet[last_column][1:]:
        cell.number_format = "#,##0"

    if sheet.max_row > 1:
        table = Table(
            displayName="NameAliasTable",
            ref=f"A1:{last_column}{sheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)


def export_data(
    data: NormalizedData,
    output_dir: Path,
    export_csv: bool = True,
    excel: bool = True,
    parquet: bool = True,
    ssb: bool = False,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    tables = data.tables()
    schemas = data.schemas()
    frames: dict[str, pd.DataFrame] = {}
    csv_dir = output_dir / "csv"
    excel_dir = output_dir / "excel"
    parquet_dir = output_dir / "parquet"
    if export_csv:
        csv_dir.mkdir(parents=True, exist_ok=True)
    if excel:
        excel_dir.mkdir(parents=True, exist_ok=True)
    if parquet:
        parquet_dir.mkdir(parents=True, exist_ok=True)

    for name, rows in tables.items():
        frame = pd.DataFrame(rows, columns=schemas[name])
        frames[name] = frame
        if export_csv:
            frame.to_csv(
                csv_dir / f"{name}.csv",
                index=False,
                quoting=csv.QUOTE_MINIMAL,
                encoding="utf-8",
            )
    if excel:
        with pd.ExcelWriter(excel_dir / "sanctions.xlsx", engine="openpyxl") as writer:
            for name, frame in frames.items():
                frame.to_excel(writer, sheet_name=name[:31], index=False)
            name_alias_frame = _name_alias_frame(data)
            name_alias_frame.to_excel(
                writer,
                sheet_name=NAME_ALIAS_SHEET,
                index=False,
            )
            _format_name_alias_sheet(writer.book[NAME_ALIAS_SHEET])
    if parquet:
        for name, frame in frames.items():
            frame.to_parquet(parquet_dir / f"{name}.parquet", index=False)
    if ssb:
        export_ssb(data, output_dir / "ssb")
