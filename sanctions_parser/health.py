from __future__ import annotations

import csv
import json
import time
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import pyarrow.parquet as pq
import requests
from lxml import etree
from openpyxl import load_workbook

from .config import SourceConfig
from .downloader import _session, sha256_file, validate_xml
from .pipeline import PARSER_SCHEMA_VERSION

PROBE_LIMIT_BYTES = 64 * 1024


@dataclass(frozen=True)
class UplinkHealth:
    state: str
    http_status: int | None
    latency_seconds: float
    redirects: int
    xml_root: str
    final_url: str
    detail: str
    source_date: datetime | None = None


@dataclass(frozen=True)
class LocalHealth:
    raw_state: str
    raw_detail: str
    parser_state: str
    parser_detail: str
    output_state: str
    output_detail: str
    last_run: str
    issues: tuple[str, ...]
    source_date: datetime | None = None


@dataclass(frozen=True)
class ProviderHealth:
    source: str
    uplink: UplinkHealth
    local: LocalHealth


@dataclass(frozen=True)
class Freshness:
    state: str
    label: str
    detail: str


def _human_size(size: int) -> str:
    value = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.1f} {unit}" if unit != "B" else f"{int(value)} B"
        value /= 1024
    return f"{value:.1f} GB"


def _parse_source_date(source_name: str, value: str) -> datetime | None:
    cleaned = value.strip()
    if not cleaned:
        return None
    try:
        if source_name == "ofac":
            return datetime.strptime(cleaned, "%m/%d/%Y").replace(tzinfo=UTC)
        elif source_name == "uk":
            return datetime.strptime(cleaned, "%d/%m/%Y").replace(tzinfo=UTC)
        parsed = datetime.fromisoformat(cleaned)
    except ValueError:
        return None
    return parsed.replace(tzinfo=UTC) if parsed.tzinfo is None else parsed


def _xml_prefix_metadata(
    source_name: str,
    chunks: Iterable[bytes],
) -> tuple[str, datetime | None, int]:
    parser = etree.XMLPullParser(events=("start", "end"), recover=True)
    root_name = ""
    source_date: datetime | None = None
    bytes_read = 0
    for chunk in chunks:
        if not chunk:
            continue
        remaining = PROBE_LIMIT_BYTES - bytes_read
        if remaining <= 0:
            break
        fragment = chunk[:remaining]
        bytes_read += len(fragment)
        parser.feed(fragment)
        for event_name, element in parser.read_events():
            element_name = etree.QName(element).localname
            if event_name == "start" and not root_name:
                root_name = element_name
                for attribute in ("dateGenerated", "generationDate"):
                    value = element.get(attribute)
                    if value:
                        source_date = _parse_source_date(source_name, value)
                        break
            elif event_name == "end" and element_name in {
                "Publish_Date",
                "DateGenerated",
            }:
                source_date = _parse_source_date(source_name, element.text or "")
            if root_name and source_date is not None:
                return root_name, source_date, bytes_read
        if len(fragment) < len(chunk):
            break
    return root_name, source_date, bytes_read


def read_source_date(path: Path, source_name: str) -> datetime | None:
    with path.open("rb") as handle:
        chunks = iter(lambda: handle.read(8192), b"")
        _root_name, source_date, _bytes_read = _xml_prefix_metadata(source_name, chunks)
    return source_date


def compare_source_dates(
    local_date: datetime | None,
    live_date: datetime | None,
) -> Freshness:
    if live_date is None:
        return Freshness(
            "unknown",
            "UNKNOWN",
            "The provider XML date could not be read",
        )
    if local_date is None:
        return Freshness(
            "unknown",
            "LOCAL DATE UNKNOWN",
            "The cached XML date could not be read",
        )
    if live_date > local_date:
        return Freshness(
            "update",
            "UPDATE AVAILABLE",
            "The provider has published a newer dated XML file",
        )
    if live_date < local_date:
        return Freshness(
            "ahead",
            "LOCAL AHEAD",
            "The provider is serving an older dated XML file",
        )
    return Freshness(
        "current",
        "NO NEWER DATE",
        "The cached and provider XML dates match",
    )


def check_uplink(source: SourceConfig) -> UplinkHealth:
    """Probe the HTTP endpoint and XML root without downloading the full file."""
    started = time.monotonic()
    session = _session(min(source.retries, 2))
    status_code: int | None = None
    redirects = 0
    final_url = source.url
    try:
        with session.get(
            source.url,
            stream=True,
            timeout=(5, min(source.timeout_seconds, 15)),
            headers={"User-Agent": source.user_agent, "Accept": "application/xml,*/*"},
            allow_redirects=True,
        ) as response:
            status_code = response.status_code
            redirects = len(response.history)
            final_url = response.url or source.url
            if not 200 <= status_code < 300:
                return UplinkHealth(
                    "degraded",
                    status_code,
                    time.monotonic() - started,
                    redirects,
                    "",
                    final_url,
                    f"Provider returned HTTP {status_code}",
                )

            root_name, source_date, bytes_read = _xml_prefix_metadata(
                source.name,
                response.iter_content(chunk_size=8192),
            )

            latency = time.monotonic() - started
            if not root_name:
                return UplinkHealth(
                    "degraded",
                    status_code,
                    latency,
                    redirects,
                    "",
                    final_url,
                    f"No XML root detected in the first {_human_size(bytes_read)}",
                )
            if root_name.lower() == "error":
                return UplinkHealth(
                    "degraded",
                    status_code,
                    latency,
                    redirects,
                    root_name,
                    final_url,
                    "Provider returned an XML error document",
                )
            if source.expected_root and root_name != source.expected_root:
                return UplinkHealth(
                    "degraded",
                    status_code,
                    latency,
                    redirects,
                    root_name,
                    final_url,
                    f"Expected XML root {source.expected_root!r}",
                )
            return UplinkHealth(
                "healthy",
                status_code,
                latency,
                redirects,
                root_name,
                final_url,
                "Endpoint and XML payload verified",
                source_date,
            )
    except (requests.RequestException, etree.XMLSyntaxError) as exc:
        return UplinkHealth(
            "failed",
            status_code,
            time.monotonic() - started,
            redirects,
            "",
            final_url,
            str(exc),
        )
    finally:
        session.close()


def _latest_current_output(
    project_root: Path, source_name: str, checksum: str
) -> tuple[Path, dict[str, object]] | None:
    output_root = project_root / "output" / source_name
    if not output_root.exists():
        return None
    candidates = sorted(
        (item for item in output_root.iterdir() if item.is_dir()),
        key=lambda item: item.name,
        reverse=True,
    )
    for output_dir in candidates:
        try:
            manifest = json.loads(
                (output_dir / "manifest.json").read_text(encoding="utf-8")
            )
        except (OSError, json.JSONDecodeError):
            continue
        if (
            manifest.get("checksum") == checksum
            and manifest.get("parser_schema_version") == PARSER_SCHEMA_VERSION
        ):
            return output_dir, manifest
    return None


def _missing_export_files(output_dir: Path, manifest: dict[str, object]) -> list[str]:
    tables = set(dict(manifest.get("tables", {})))
    formats = set(manifest.get("formats", []))
    missing: list[str] = []
    if "excel" in formats and not (output_dir / "excel" / "sanctions.xlsx").is_file():
        missing.append("excel/sanctions.xlsx")
    if "csv" in formats:
        missing.extend(
            f"csv/{table}.csv"
            for table in tables
            if not (output_dir / "csv" / f"{table}.csv").is_file()
        )
    if "parquet" in formats:
        missing.extend(
            f"parquet/{table}.parquet"
            for table in tables
            if not (output_dir / "parquet" / f"{table}.parquet").is_file()
        )
    if "ssb" in formats:
        source = str(manifest.get("source", "")).casefold()
        missing.extend(
            f"ssb/{source}-{party_type}.csv"
            for party_type in ("individuals", "organizations")
            if not (output_dir / "ssb" / f"{source}-{party_type}.csv").is_file()
        )
    return sorted(missing)


def _blank_primary_names(output_dir: Path) -> int:
    parquet = output_dir / "parquet" / "entity.parquet"
    if parquet.is_file():
        values = pq.read_table(parquet, columns=["primary_name"]).column(0).to_pylist()
        return sum(not str(value or "").strip() for value in values)

    csv_path = output_dir / "csv" / "entity.csv"
    if csv_path.is_file():
        with csv_path.open("r", encoding="utf-8", newline="") as handle:
            return sum(
                not str(row.get("primary_name", "")).strip()
                for row in csv.DictReader(handle)
            )

    workbook_path = output_dir / "excel" / "sanctions.xlsx"
    if workbook_path.is_file():
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            sheet = workbook["entity"]
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            column = headers.index("primary_name") + 1
            return sum(
                not str(sheet.cell(row, column).value or "").strip()
                for row in range(2, sheet.max_row + 1)
            )
        finally:
            workbook.close()
    raise FileNotFoundError("No entity export is available")


def check_local_health(source: SourceConfig, project_root: Path) -> LocalHealth:
    """Validate cached raw state and the latest matching normalized output."""
    issues: list[str] = []
    state_path = project_root / ".state" / f"{source.name}.json"
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
        raw_path = Path(state["path"])
        expected_checksum = str(state["checksum"])
    except (OSError, KeyError, json.JSONDecodeError) as exc:
        message = f"Raw state is missing or invalid: {exc}"
        return LocalHealth(
            "failed",
            "State unavailable",
            "failed",
            "Not checked",
            "failed",
            "Not checked",
            "Never",
            (message,),
        )

    if not raw_path.is_file():
        message = f"Cached raw XML is missing: {raw_path}"
        return LocalHealth(
            "failed",
            "File missing",
            "failed",
            "Not checked",
            "failed",
            "Not checked",
            "Never",
            (message,),
        )

    source_date: datetime | None = None
    try:
        validate_xml(raw_path, source.expected_root)
        actual_checksum = sha256_file(raw_path, source.chunk_size)
        source_date = read_source_date(raw_path, source.name)
    except (OSError, ValueError) as exc:
        issues.append(f"Raw XML failed validation: {exc}")
        raw_state = "failed"
        raw_detail = "Invalid XML"
    else:
        if actual_checksum != expected_checksum:
            issues.append("Raw XML checksum does not match its state record")
            raw_state = "failed"
            raw_detail = "Checksum mismatch"
        else:
            raw_state = "healthy"
            raw_detail = f"Valid · {_human_size(raw_path.stat().st_size)}"

    current = _latest_current_output(project_root, source.name, expected_checksum)
    if current is None:
        issues.append("No current export matches the raw checksum and parser version")
        return LocalHealth(
            raw_state,
            raw_detail,
            "warning",
            "Rebuild needed",
            "warning",
            "Current export missing",
            "Never",
            tuple(issues),
            source_date,
        )

    output_dir, manifest = current
    table_counts = dict(manifest.get("tables", {}))
    entity_count = int(table_counts.get("entity", 0))
    try:
        blank_names = _blank_primary_names(output_dir)
    except (OSError, KeyError, ValueError) as exc:
        blank_names = -1
        issues.append(f"Could not inspect primary names: {exc}")

    if entity_count <= 0:
        parser_state = "failed"
        parser_detail = "Zero entities"
        issues.append("The current parser output contains zero entities")
    elif blank_names > 0:
        parser_state = "warning"
        parser_detail = f"{entity_count:,} · {blank_names:,} blank names"
        issues.append(f"{blank_names:,} entities have blank primary names")
    elif blank_names == 0:
        parser_state = "healthy"
        parser_detail = f"{entity_count:,} entities"
    else:
        parser_state = "warning"
        parser_detail = f"{entity_count:,} · names unchecked"

    missing = _missing_export_files(output_dir, manifest)
    if missing:
        output_state = "failed"
        output_detail = f"{len(missing)} files missing"
        issues.append(
            "Missing export files: "
            + ", ".join(missing[:3])
            + ("…" if len(missing) > 3 else "")
        )
    else:
        formats = set(manifest.get("formats", []))
        labels = [
            label
            for key, label in (
                ("csv", "CSV"),
                ("excel", "XLSX"),
                ("parquet", "PQ"),
                ("ssb", "SSB"),
            )
            if key in formats
        ]
        output_state = "healthy"
        output_detail = f"{len(table_counts)} tables · {'/'.join(labels)}"

    modified = (
        datetime.fromtimestamp(
            (output_dir / "manifest.json").stat().st_mtime,
            tz=UTC,
        )
        .astimezone()
        .strftime("%Y-%m-%d %H:%M")
    )
    return LocalHealth(
        raw_state,
        raw_detail,
        parser_state,
        parser_detail,
        output_state,
        output_detail,
        modified,
        tuple(issues),
        source_date,
    )


def check_provider_health(source: SourceConfig, project_root: Path) -> ProviderHealth:
    """Run the live and local checks for a single provider."""
    return ProviderHealth(
        source.name,
        check_uplink(source),
        check_local_health(source, project_root),
    )
