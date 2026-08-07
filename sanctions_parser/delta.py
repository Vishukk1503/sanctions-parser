from __future__ import annotations

import csv
import gzip
import json
import re
import unicodedata
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .models import NormalizedData
from .reporting import alias_tag

DELTA_SCHEMA_VERSION = 1
CHANGE_COLUMNS = [
    "status",
    "entity_id",
    "source",
    "record_type",
    "primary_name",
    "aliases",
    "alias_count",
    "what_changed",
    "previous_primary_name",
    "aliases_added",
    "aliases_removed",
    "provider_date_updated",
]
SUMMARY_COLUMNS = [
    "source",
    "result",
    "previous_checkpoint",
    "current_check",
    "new",
    "updated",
    "removed",
    "unchanged",
    "current_records",
    "previous_checksum",
    "current_checksum",
]
ALIAS_PRIORITY = {"strong": 0, "weak": 1, "former": 2}
RELATION_LABELS = {
    "address": "Addresses",
    "document": "Documents",
    "nationality": "Nationalities",
    "program": "Programs",
    "relationship": "Relationships",
    "date_of_birth": "Birth dates",
    "place_of_birth": "Birth places",
    "designation": "Designations",
    "regulation": "Regulations",
    "contact": "Contacts",
    "sanction": "Sanctions",
}
ENTITY_FIELD_LABELS = {
    "record_type": "Record type",
    "status": "Status",
    "comments": "Comments",
    "date_listed": "Listed date",
    "date_updated": "Provider update date",
}


@dataclass(frozen=True)
class DeltaSnapshot:
    source: str
    checksum: str
    parser_schema_version: int
    raw_path: str
    created_at: str
    tables: dict[str, list[dict[str, Any]]]

    @property
    def entity_count(self) -> int:
        return len(self.tables.get("entity", []))


@dataclass(frozen=True)
class DeltaSummary:
    source: str
    result: str
    previous_checkpoint: str
    current_check: str
    new: int
    updated: int
    removed: int
    unchanged: int
    current_records: int
    previous_checksum: str
    current_checksum: str

    @property
    def changed(self) -> int:
        return self.new + self.updated + self.removed

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, value: object) -> DeltaSummary | None:
        if not isinstance(value, dict):
            return None
        try:
            return cls(
                source=str(value["source"]),
                result=str(value["result"]),
                previous_checkpoint=str(value.get("previous_checkpoint", "")),
                current_check=str(value["current_check"]),
                new=int(value["new"]),
                updated=int(value["updated"]),
                removed=int(value["removed"]),
                unchanged=int(value["unchanged"]),
                current_records=int(value["current_records"]),
                previous_checksum=str(value.get("previous_checksum", "")),
                current_checksum=str(value["current_checksum"]),
            )
        except (KeyError, TypeError, ValueError):
            return None


@dataclass(frozen=True)
class DeltaComparison:
    summary: DeltaSummary
    changes: tuple[dict[str, Any], ...]


@dataclass(frozen=True)
class DeltaBaselineRef:
    source: str
    checksum: str
    parser_schema_version: int
    raw_path: Path
    snapshot_path: Path
    created_at: str
    last_checked_at: str


@dataclass(frozen=True)
class StoredDeltaReport:
    source: str
    summary: DeltaSummary
    output_dir: Path
    changes_path: Path | None
    warnings: tuple[str, ...] = ()


def _text(value: object) -> str:
    return " ".join(unicodedata.normalize("NFC", str(value or "")).split())


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".part")
    temporary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    temporary.replace(path)


def build_snapshot(
    data: NormalizedData,
    *,
    source: str,
    checksum: str,
    parser_schema_version: int,
    raw_path: Path,
    created_at: str | None = None,
) -> DeltaSnapshot:
    return DeltaSnapshot(
        source=source,
        checksum=checksum,
        parser_schema_version=parser_schema_version,
        raw_path=str(raw_path),
        created_at=created_at or datetime.now(UTC).isoformat(),
        tables=data.tables(),
    )


def save_snapshot(snapshot: DeltaSnapshot, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".part")
    with gzip.open(temporary, "wt", encoding="utf-8") as handle:
        json.dump(asdict(snapshot), handle, ensure_ascii=False, separators=(",", ":"))
    temporary.replace(path)


def load_snapshot(path: Path) -> DeltaSnapshot:
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict) or not isinstance(value.get("tables"), dict):
        raise TypeError(f"Invalid delta snapshot: {path}")
    return DeltaSnapshot(
        source=str(value["source"]),
        checksum=str(value["checksum"]),
        parser_schema_version=int(value["parser_schema_version"]),
        raw_path=str(value["raw_path"]),
        created_at=str(value["created_at"]),
        tables=value["tables"],
    )


def delta_state_path(project_root: Path, source: str) -> Path:
    return project_root / ".state" / "delta" / f"{source}.json"


def load_baseline_ref(project_root: Path, source: str) -> DeltaBaselineRef | None:
    path = delta_state_path(project_root, source)
    if not path.exists():
        return None
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return DeltaBaselineRef(
            source=source,
            checksum=str(value["checksum"]),
            parser_schema_version=int(value["parser_schema_version"]),
            raw_path=Path(value["raw_path"]),
            snapshot_path=Path(value["snapshot_path"]),
            created_at=str(value["created_at"]),
            last_checked_at=str(value.get("last_checked_at", value["created_at"])),
        )
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError(
            f"Invalid {source.upper()} delta baseline state: {path}"
        ) from exc


def save_baseline(
    project_root: Path,
    snapshot: DeltaSnapshot,
    *,
    last_checked_at: str,
) -> DeltaBaselineRef:
    safe_stamp = re.sub(r"[^0-9]", "", snapshot.created_at)[:20]
    snapshot_path = (
        project_root
        / ".state"
        / "delta"
        / "snapshots"
        / snapshot.source
        / f"{safe_stamp}-{snapshot.checksum[:12]}.json.gz"
    )
    if not snapshot_path.exists():
        save_snapshot(snapshot, snapshot_path)
    payload = {
        "source": snapshot.source,
        "checksum": snapshot.checksum,
        "parser_schema_version": snapshot.parser_schema_version,
        "raw_path": snapshot.raw_path,
        "snapshot_path": str(snapshot_path),
        "created_at": snapshot.created_at,
        "last_checked_at": last_checked_at,
        "entity_count": snapshot.entity_count,
    }
    _atomic_json(delta_state_path(project_root, snapshot.source), payload)
    return DeltaBaselineRef(
        source=snapshot.source,
        checksum=snapshot.checksum,
        parser_schema_version=snapshot.parser_schema_version,
        raw_path=Path(snapshot.raw_path),
        snapshot_path=snapshot_path,
        created_at=snapshot.created_at,
        last_checked_at=last_checked_at,
    )


def refresh_unchanged_baseline(
    project_root: Path,
    baseline: DeltaBaselineRef,
    *,
    raw_path: Path,
    checked_at: str,
) -> None:
    """Advance the check time without storing a duplicate unchanged snapshot."""
    payload = {
        "source": baseline.source,
        "checksum": baseline.checksum,
        "parser_schema_version": baseline.parser_schema_version,
        "raw_path": str(raw_path),
        "snapshot_path": str(baseline.snapshot_path),
        "created_at": checked_at,
        "last_checked_at": checked_at,
    }
    _atomic_json(delta_state_path(project_root, baseline.source), payload)


def _entity_rows(snapshot: DeltaSnapshot) -> dict[str, dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    for row in snapshot.tables.get("entity", []):
        entity_id = _text(row.get("entity_id"))
        if not entity_id:
            continue
        if entity_id in rows:
            raise ValueError(
                f"Duplicate {snapshot.source} ID in delta snapshot: {entity_id}"
            )
        rows[entity_id] = row
    return rows


def _rows_by_entity(
    snapshot: DeltaSnapshot,
    table_name: str,
) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in snapshot.tables.get(table_name, []):
        entity_id = _text(row.get("entity_id"))
        if entity_id:
            result[entity_id].append(row)
    return result


def _canonical_row(row: dict[str, Any], table_name: str) -> tuple[tuple[str, str], ...]:
    values: list[tuple[str, str]] = []
    for key, value in row.items():
        if key == "entity_id":
            continue
        normalized = _text(value)
        if table_name == "alias" and key == "quality":
            normalized = alias_tag(normalized)
        values.append((key, normalized))
    return tuple(sorted(values))


def _all_relation_sets(
    snapshot: DeltaSnapshot,
) -> dict[str, dict[str, frozenset[tuple[tuple[str, str], ...]]]]:
    result: dict[str, dict[str, frozenset[tuple[tuple[str, str], ...]]]] = {}
    for table_name in RELATION_LABELS:
        grouped = _rows_by_entity(snapshot, table_name)
        result[table_name] = {
            entity_id: frozenset(_canonical_row(row, table_name) for row in rows)
            for entity_id, rows in grouped.items()
        }
    return result


def _all_alias_maps(
    snapshot: DeltaSnapshot,
    entities: dict[str, dict[str, Any]],
) -> dict[str, dict[tuple[str, str], dict[str, str]]]:
    result: dict[str, dict[tuple[str, str], dict[str, str]]] = defaultdict(dict)
    for entity_id, rows in _rows_by_entity(snapshot, "alias").items():
        primary_key = _text(entities.get(entity_id, {}).get("primary_name")).casefold()
        aliases = result[entity_id]
        for row in rows:
            value = _text(row.get("alias"))
            if not value or value.casefold() == primary_key:
                continue
            language = _text(row.get("language"))
            category = alias_tag(_text(row.get("quality")))
            key = (value.casefold(), language.casefold())
            candidate = {
                "value": value,
                "category": category,
                "language": language,
            }
            existing = aliases.get(key)
            if (
                existing is None
                or ALIAS_PRIORITY[category] < ALIAS_PRIORITY[existing["category"]]
            ):
                aliases[key] = candidate
    return dict(result)


def _alias_display(values: list[dict[str, str]]) -> str:
    ordered = sorted(
        values,
        key=lambda item: (
            ALIAS_PRIORITY[item["category"]],
            item["value"].casefold(),
            item["language"].casefold(),
        ),
    )
    return " | ".join(f"{item['value']} [{item['category']}]" for item in ordered)


def _changed_details(
    entity_id: str,
    previous_entity: dict[str, Any],
    current_entity: dict[str, Any],
    old_aliases: dict[tuple[str, str], dict[str, str]],
    new_aliases: dict[tuple[str, str], dict[str, str]],
    previous_relations: dict[str, dict[str, frozenset[tuple[tuple[str, str], ...]]]],
    current_relations: dict[str, dict[str, frozenset[tuple[tuple[str, str], ...]]]],
) -> tuple[list[str], list[dict[str, str]], list[dict[str, str]]]:
    details: list[str] = []
    if _text(previous_entity.get("primary_name")) != _text(
        current_entity.get("primary_name")
    ):
        details.append("Primary name changed")

    added: list[dict[str, str]] = []
    removed: list[dict[str, str]] = []
    strength_changes = 0
    spelling_changes = 0
    for key in sorted(old_aliases.keys() | new_aliases.keys()):
        old = old_aliases.get(key)
        new = new_aliases.get(key)
        if old is None and new is not None:
            added.append(new)
        elif new is None and old is not None:
            removed.append(old)
        elif old is not None and new is not None:
            if old["category"] != new["category"]:
                strength_changes += 1
            if old["value"] != new["value"] or old["language"] != new["language"]:
                spelling_changes += 1
    if added:
        details.append(f"{len(added)} alias{'es' if len(added) != 1 else ''} added")
    if removed:
        details.append(
            f"{len(removed)} alias{'es' if len(removed) != 1 else ''} removed"
        )
    if strength_changes:
        details.append(
            f"{strength_changes} alias strength change"
            f"{'s' if strength_changes != 1 else ''}"
        )
    if spelling_changes:
        details.append(
            f"{spelling_changes} alias text/language change"
            f"{'s' if spelling_changes != 1 else ''}"
        )

    for field, label in ENTITY_FIELD_LABELS.items():
        if _text(previous_entity.get(field)) != _text(current_entity.get(field)):
            details.append(f"{label} changed")
    for table_name, label in RELATION_LABELS.items():
        if previous_relations[table_name].get(
            entity_id, frozenset()
        ) != current_relations[table_name].get(entity_id, frozenset()):
            details.append(f"{label} changed")
    return details, added, removed


def compare_snapshots(
    previous: DeltaSnapshot | None,
    current: DeltaSnapshot,
    *,
    checked_at: str | None = None,
) -> DeltaComparison:
    now = checked_at or datetime.now(UTC).isoformat()
    current_entities = _entity_rows(current)
    if not current_entities:
        raise ValueError(
            f"{current.source.upper()} delta snapshot contains zero entities"
        )

    if previous is None:
        return DeltaComparison(
            DeltaSummary(
                current.source,
                "baseline_created",
                "",
                now,
                0,
                0,
                0,
                current.entity_count,
                current.entity_count,
                "",
                current.checksum,
            ),
            (),
        )

    if previous.source != current.source:
        raise ValueError("Delta snapshots belong to different providers")
    previous_entities = _entity_rows(previous)
    previous_aliases = _all_alias_maps(previous, previous_entities)
    current_aliases = _all_alias_maps(current, current_entities)
    previous_relations = _all_relation_sets(previous)
    current_relations = _all_relation_sets(current)
    previous_ids = set(previous_entities)
    current_ids = set(current_entities)
    added_ids = current_ids - previous_ids
    removed_ids = previous_ids - current_ids
    common_ids = previous_ids & current_ids
    rows: list[dict[str, Any]] = []
    updated = 0

    def change_row(
        status: str,
        entity_id: str,
        display_snapshot: DeltaSnapshot,
        display_entity: dict[str, Any],
        what_changed: str,
        previous_name: str = "",
        aliases_added: list[dict[str, str]] | None = None,
        aliases_removed: list[dict[str, str]] | None = None,
    ) -> dict[str, Any]:
        alias_index = (
            current_aliases if display_snapshot is current else previous_aliases
        )
        aliases = list(alias_index.get(entity_id, {}).values())
        return {
            "status": status,
            "entity_id": entity_id,
            "source": current.source,
            "record_type": _text(display_entity.get("record_type")),
            "primary_name": _text(display_entity.get("primary_name")),
            "aliases": _alias_display(aliases),
            "alias_count": len(aliases),
            "what_changed": what_changed,
            "previous_primary_name": previous_name,
            "aliases_added": _alias_display(aliases_added or []),
            "aliases_removed": _alias_display(aliases_removed or []),
            "provider_date_updated": _text(display_entity.get("date_updated")),
        }

    for entity_id in sorted(added_ids):
        rows.append(
            change_row(
                "NEW",
                entity_id,
                current,
                current_entities[entity_id],
                "New record",
            )
        )
    for entity_id in sorted(common_ids):
        old_entity = previous_entities[entity_id]
        new_entity = current_entities[entity_id]
        details, aliases_added, aliases_removed = _changed_details(
            entity_id,
            old_entity,
            new_entity,
            previous_aliases.get(entity_id, {}),
            current_aliases.get(entity_id, {}),
            previous_relations,
            current_relations,
        )
        if not details:
            continue
        updated += 1
        rows.append(
            change_row(
                "UPDATED",
                entity_id,
                current,
                new_entity,
                "; ".join(details),
                previous_name=_text(old_entity.get("primary_name")),
                aliases_added=aliases_added,
                aliases_removed=aliases_removed,
            )
        )
    for entity_id in sorted(removed_ids):
        old_entity = previous_entities[entity_id]
        rows.append(
            change_row(
                "REMOVED",
                entity_id,
                previous,
                old_entity,
                "Removed from current source",
                previous_name=_text(old_entity.get("primary_name")),
            )
        )

    unchanged = len(common_ids) - updated
    changed_count = len(added_ids) + updated + len(removed_ids)
    return DeltaComparison(
        DeltaSummary(
            current.source,
            "changes" if changed_count else "no_changes",
            previous.created_at,
            now,
            len(added_ids),
            updated,
            len(removed_ids),
            unchanged,
            current.entity_count,
            previous.checksum,
            current.checksum,
        ),
        tuple(rows),
    )


def _format_sheet(sheet: Worksheet, widths: dict[str, int], table_name: str) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="0F4C81")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)


def export_delta_report(comparison: DeltaComparison, output_dir: Path) -> None:
    """Write the compact delta report in CSV, Excel and Parquet formats."""
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_dir = output_dir / "csv"
    excel_dir = output_dir / "excel"
    parquet_dir = output_dir / "parquet"
    for directory in (csv_dir, excel_dir, parquet_dir):
        directory.mkdir(parents=True, exist_ok=True)

    summary_frame = pd.DataFrame(
        [comparison.summary.to_dict()], columns=SUMMARY_COLUMNS
    )
    changes_frame = pd.DataFrame(comparison.changes, columns=CHANGE_COLUMNS)
    summary_frame.to_csv(
        csv_dir / "change_summary.csv",
        index=False,
        quoting=csv.QUOTE_MINIMAL,
        encoding="utf-8",
    )
    changes_frame.to_csv(
        csv_dir / "name_alias_changes.csv",
        index=False,
        quoting=csv.QUOTE_MINIMAL,
        encoding="utf-8",
    )
    summary_frame.to_parquet(parquet_dir / "change_summary.parquet", index=False)
    changes_frame.to_parquet(parquet_dir / "name_alias_changes.parquet", index=False)

    workbook_path = excel_dir / "delta-report.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        summary_frame.to_excel(writer, sheet_name="change summary", index=False)
        changes_frame.to_excel(writer, sheet_name="name+alias changes", index=False)
        summary_sheet = writer.book["change summary"]
        changes_sheet = writer.book["name+alias changes"]
        _format_sheet(
            summary_sheet,
            {
                "A": 12,
                "B": 18,
                "C": 24,
                "D": 24,
                "E": 10,
                "F": 10,
                "G": 10,
                "H": 12,
                "I": 16,
                "J": 18,
                "K": 18,
            },
            "DeltaSummaryTable",
        )
        _format_sheet(
            changes_sheet,
            {
                "A": 12,
                "B": 18,
                "C": 12,
                "D": 16,
                "E": 38,
                "F": 70,
                "G": 12,
                "H": 50,
                "I": 38,
                "J": 55,
                "K": 55,
                "L": 22,
            },
            "NameAliasChangesTable",
        )
        row_fills = {
            "NEW": PatternFill("solid", fgColor="E2F0D9"),
            "UPDATED": PatternFill("solid", fgColor="FCE4D6"),
            "REMOVED": PatternFill("solid", fgColor="F4CCCC"),
        }
        for row in changes_sheet.iter_rows(min_row=2):
            fill = row_fills.get(str(row[0].value or ""))
            if fill:
                for cell in row:
                    cell.fill = fill


def load_latest_delta_report(
    project_root: Path,
    source: str,
) -> StoredDeltaReport | None:
    root = project_root / "output" / "delta" / source
    if not root.exists():
        return None
    candidates = sorted(
        (path for path in root.iterdir() if path.is_dir()),
        key=lambda path: path.name,
        reverse=True,
    )
    for output_dir in candidates:
        try:
            manifest = json.loads(
                (output_dir / "manifest.json").read_text(encoding="utf-8")
            )
        except (OSError, json.JSONDecodeError):
            continue
        summary = DeltaSummary.from_dict(manifest.get("summary"))
        if summary is None:
            continue
        changes_path = output_dir / "csv" / "name_alias_changes.csv"
        warnings = manifest.get("warnings", [])
        return StoredDeltaReport(
            source=source,
            summary=summary,
            output_dir=output_dir,
            changes_path=changes_path if changes_path.is_file() else None,
            warnings=(
                tuple(str(item) for item in warnings)
                if isinstance(warnings, list)
                else ()
            ),
        )
    return None


def load_delta_preview(
    report: StoredDeltaReport, limit: int = 10
) -> list[dict[str, str]]:
    if report.changes_path is None:
        return []
    with report.changes_path.open("r", encoding="utf-8", newline="") as handle:
        return list(next_rows(csv.DictReader(handle), limit))


def next_rows(reader: csv.DictReader, limit: int):
    for index, row in enumerate(reader):
        if index >= limit:
            break
        yield dict(row)
